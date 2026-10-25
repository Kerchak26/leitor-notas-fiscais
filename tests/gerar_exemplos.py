"""Gera PDFs fictícios de notas fiscais e boletos para testes e demonstração.

Uso: python tests/gerar_exemplos.py [pasta_destino]
"""

import sys
from datetime import date
from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

# Uma linha pode ser um texto simples ou uma lista de (x, texto) para montar colunas
Linha = str | list[tuple[int, str]]


# ------------------------------------------------- dígitos verificadores ----

def _mod11(numero: str) -> int:
    pesos = [2, 3, 4, 5, 6, 7, 8, 9]
    return sum(int(d) * pesos[i % 8] for i, d in enumerate(reversed(numero))) % 11


def chave_ficticia(cnpj: str, serie: int, numero: int) -> str:
    base = f"35" "2610" f"{cnpj}" "55" f"{serie:03d}" f"{numero:09d}" "1" "12345678"
    resto = _mod11(base)
    return base + str(0 if resto < 2 else 11 - resto)


def _mod10(numero: str) -> str:
    soma = 0
    for i, d in enumerate(reversed(numero)):
        p = int(d) * (2 if i % 2 == 0 else 1)
        soma += p // 10 + p % 10
    return str((10 - soma % 10) % 10)


def linha_digitavel(vencimento: date, valor_centavos: int, livre: str, banco: str = "237") -> str:
    """Monta uma linha digitável válida (padrão FEBRABAN) já formatada."""
    fator = 1000 + (vencimento - date(2025, 2, 22)).days  # ciclo iniciado em 22/02/2025
    fator_valor = f"{fator:04d}{valor_centavos:010d}"
    sem_dv = banco + "9" + fator_valor + livre
    dv = 11 - _mod11(sem_dv)
    dv = 1 if dv in (0, 10, 11) else dv
    c1 = banco + "9" + livre[:5]
    c2, c3 = livre[5:15], livre[15:25]
    c1, c2, c3 = c1 + _mod10(c1), c2 + _mod10(c2), c3 + _mod10(c3)
    return f"{c1[:5]}.{c1[5:]} {c2[:5]}.{c2[5:]} {c3[:5]}.{c3[5:]} {dv} {fator_valor}"


# --------------------------------------------------------------- desenho ----

def _desenhar_pagina(c: canvas.Canvas, linhas: list[Linha]) -> None:
    y = 800
    c.setFont("Helvetica", 8)
    for linha in linhas:
        for x, texto in ([(40, linha)] if isinstance(linha, str) else linha):
            c.drawString(x, y, texto)
        y -= 15
    c.showPage()


def _escrever(caminho: Path, *paginas: list[Linha]) -> Path:
    c = canvas.Canvas(str(caminho), pagesize=A4)
    for pagina in paginas:
        _desenhar_pagina(c, pagina)
    c.save()
    return caminho


# ------------------------------------------------------------ documentos ----

def pagina_danfe(numero: int = 4587, cnpj: str = "12345678000195",
                 empresa: str = "DISTRIBUIDORA ALFA COMÉRCIO LTDA", valor: str = "3.250,90",
                 vencimento: str = "13/11/2026", destinatario: str = "EMPRESA CLIENTE S.A.",
                 cnpj_destinatario: str = "98.765.432/0001-10",
                 endereco: str = "RUA DAS ACÁCIAS, 100") -> list[Linha]:
    chave = chave_ficticia(cnpj, 1, numero)
    cnpj_fmt = f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}"
    return [
        f"RECEBEMOS DE {empresa} OS PRODUTOS CONSTANTES DA NOTA FISCAL INDICADA AO LADO",
        f"NF-e Nº {numero:09,d} SÉRIE 001".replace(",", "."),
        empresa,
        "Rua das Flores, 100 - São Paulo/SP",
        "DANFE - DOCUMENTO AUXILIAR DA NOTA FISCAL ELETRÔNICA",
        "CHAVE DE ACESSO",
        " ".join(chave[i:i + 4] for i in range(0, 44, 4)),
        f"CNPJ {cnpj_fmt}    INSCRIÇÃO ESTADUAL 123.456.789.110",
        "DESTINATÁRIO / REMETENTE",
        [(40, "NOME/RAZÃO SOCIAL"), (260, "CNPJ/CPF"), (400, "DATA DA EMISSÃO")],
        [(40, destinatario), (260, cnpj_destinatario), (400, "14/10/2026")],
        [(40, "ENDEREÇO"), (260, "BAIRRO/DISTRITO"), (400, "CEP")],
        [(40, endereco), (260, "JARDIM DAS FLORES"), (400, "00000-000")],
        "FATURA / DUPLICATA",
        f"NÚM. 001   VENCIMENTO {vencimento}   VALOR {valor}",
        "CÁLCULO DO IMPOSTO",
        [(40, "VALOR DO ICMS"), (200, "VALOR TOTAL DOS PRODUTOS"), (400, "VALOR TOTAL DA NOTA")],
        [(40, "390,11"), (200, valor), (400, valor)],
    ]


def pagina_nfse() -> list[Linha]:
    return [
        "PREFEITURA MUNICIPAL DE CIDADE EXEMPLO",
        "NOTA FISCAL DE SERVIÇOS ELETRÔNICA - NFS-e",
        "Número da NFS-e: 2026000123",
        "Data de Emissão: 03/10/2026 10:42:11",
        "PRESTADOR DE SERVIÇOS",
        "Razão Social: TECH SOLUÇÕES EM SOFTWARE EIRELI",
        "CNPJ: 11.222.333/0001-81",
        "TOMADOR DE SERVIÇOS",
        "Razão Social: EMPRESA CLIENTE S.A.",
        "Discriminação: Desenvolvimento de sistema - outubro/2026",
        "Valor dos Serviços: R$ 12.800,00",
        "ISS (2%): R$ 256,00",
        "VALOR LÍQUIDO DA NOTA: R$ 12.544,00",
    ]


def pagina_boleto(beneficiario: str, cnpj: str, documento: str, vencimento: date,
                  valor: str, livre: str = "1234509000000045870012345") -> list[Linha]:
    centavos = int(valor.replace(".", "").replace(",", ""))
    linha = linha_digitavel(vencimento, centavos, livre)
    venc = vencimento.strftime("%d/%m/%Y")
    ficha = [
        [(40, "Bradesco | 237-2"), (200, linha)],
        [(40, "Local de pagamento"), (420, "Vencimento")],
        [(40, "Pagável preferencialmente na rede Bradesco ou em qualquer banco até o vencimento"), (420, venc)],
        [(40, "Beneficiário"), (420, "Agência/Código do Beneficiário")],
        [(40, f"{beneficiario} - CNPJ {cnpj}"), (420, "1234/0012345-6")],
        [(40, "Data do Documento"), (130, "Nº do Documento"), (230, "Espécie Doc."), (300, "Aceite"),
         (340, "Data Processamento"), (420, "Nosso Número")],
        [(40, "14/10/2026"), (130, documento), (230, "DM"), (300, "N"), (340, "14/10/2026"),
         (420, "09/00000004587-1")],
        [(40, "Instruções: Não receber após 30 dias do vencimento. Referente à NF " + documento),
         (420, "(=) Valor do Documento")],
        [(420, valor)],
        "Pagador: EMPRESA CLIENTE S.A. - CNPJ 98.765.432/0001-10",
        [(40, "FICHA DE COMPENSAÇÃO"), (380, "Autenticação Mecânica")],
    ]
    return [[(40, "RECIBO DO PAGADOR")], *ficha[1:], "- " * 50, *ficha]


# ------------------------------------------------------ capa de custos ----

MESES_CAPA = ["Dezembro", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
              "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]


def _aba_capa(ws, contratos: list[tuple[str, list[float | None]]], anotacoes: list[str] = ()) -> None:
    """Monta uma aba no formato das capas: Dez/2025 a Dez/2026, Total e Valor Total."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    amarelo = PatternFill("solid", fgColor="FFFF00")
    ws["C1"] = 2025
    ws["D1"] = 2026
    ws.merge_cells("D1:O1")
    ws["A2"], ws["B2"], ws["P2"] = "Item", "Contratos", "Total"
    for i, mes in enumerate(MESES_CAPA):
        ws.cell(2, 3 + i, mes)
    for linha, (nome, valores) in enumerate(contratos, start=3):
        ws.cell(linha, 1, linha - 2)
        ws.cell(linha, 2, nome)
        for i, v in enumerate(valores):
            if v is not None:
                ws.cell(linha, 3 + i, v).number_format = "#,##0.00"
        ws.cell(linha, 16, f"=SUM(C{linha}:O{linha})").number_format = "#,##0.00"
    ultima = 2 + len(contratos)
    total = ultima + 2
    ws.cell(total, 2, "Valor Total")
    for col in range(3, 17):
        letra = get_column_letter(col)
        ws.cell(total, col, f"=SUM({letra}3:{letra}{ultima + 1})").number_format = "#,##0.00"
    for linha in (1, 2, total):
        for col in range(1, 17):
            ws.cell(linha, col).fill = amarelo
            ws.cell(linha, col).font = Font(bold=True)
    # Mesmas posições do print: linhas 7, 9, 12 e 13 quando o Valor Total está na linha 5
    for deslocamento, texto in zip([2, 4, 7, 8], anotacoes):
        ws.cell(total + deslocamento, 2, texto)


def gerar_capa_exemplo(caminho: Path) -> Path:
    """Capa de custos fictícia: uma aba por fornecedor, Dez/2025 a Dez/2026."""
    from openpyxl import Workbook

    def recorrente(valor: float) -> list[float | None]:
        return [valor] * 10 + [None] * 3            # lançado de Dez/2025 a Set/2026

    wb = Workbook()
    wb.active.title = "Plan2"
    _aba_capa(wb.create_sheet("Beta Link"), [("LINK DE INTERNET DEDICADO", recorrente(6200.0))])
    _aba_capa(wb.create_sheet("Gama - Telefonia Móvel"), [("LINHAS MÓVEIS", recorrente(2300.0))])
    _aba_capa(wb.create_sheet("Alfa Telecom"), [("ALFA TELECOM", recorrente(18450.0))], [
        "Item 01 - Locação e suporte de 80 coletores de dados",
        "Contrato centralizado na Unidade Centro a partir de 2026",
        "Nota 10101 - Vencimento 14/09/2026",
        "Nota 20202 - Vencimento 13/09/2026",
    ])
    _aba_capa(wb.create_sheet("Delta Print"), [
        ("Locação de impressoras", recorrente(3900.0)),
        ("Franquia de cópias excedentes", recorrente(640.0)),
    ])
    _aba_capa(wb.create_sheet("Sigma"), [("SISTEMA DE GESTÃO", recorrente(11000.0))])
    _aba_capa(wb.create_sheet("Sigma - Assinatura Digital"), [("ASSINATURA DIGITAL", recorrente(750.0))])
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho


# ------------------------------------------------- arquivos de exemplo ----

ALFA = ("DISTRIBUIDORA ALFA COMÉRCIO LTDA", "12.345.678/0001-95")
TECH = ("TECH SOLUÇÕES EM SOFTWARE EIRELI", "11.222.333/0001-81")


def gerar_danfe(caminho: Path) -> Path:
    return _escrever(caminho, pagina_danfe())


def gerar_nfse(caminho: Path) -> Path:
    return _escrever(caminho, pagina_nfse())


def gerar_boleto_alfa(caminho: Path) -> Path:
    return _escrever(caminho, pagina_boleto(*ALFA, "4587/1", date(2026, 11, 13), "3.250,90"))


def gerar_boletos_tech(caminho: Path) -> Path:
    """Duas parcelas da NFS-e, cada uma em uma página."""
    return _escrever(
        caminho,
        pagina_boleto(*TECH, "2026000123/1", date(2026, 11, 3), "6.272,00", "9876500000000000000123401"),
        pagina_boleto(*TECH, "2026000123/2", date(2026, 12, 3), "6.272,00", "9876500000000000000123402"),
    )


def gerar_nota_e_boleto_mesmo_pdf(caminho: Path) -> Path:
    return _escrever(caminho, pagina_danfe(),
                     pagina_boleto(*ALFA, "4587/1", date(2026, 11, 13), "3.250,90"))


def gerar_notas_alfa_telecom(pasta: Path) -> None:
    """E-mail com duas notas do mesmo contrato e seus boletos (somam o valor do mês)."""
    pasta.mkdir(parents=True, exist_ok=True)
    emitente = dict(cnpj="44555666000171", empresa="ALFA TELECOM COMÉRCIO E SERVIÇOS LTDA",
                    destinatario="EMPRESA CLIENTE S.A. - UNIDADE CENTRO", cnpj_destinatario="11.111.111/0001-11")
    beneficiario = ("ALFA TELECOM COMÉRCIO E SERVIÇOS LTDA", "44.555.666/0001-71")
    for numero, valor, venc, livre in [(50001, "8.000,00", date(2026, 10, 14), "1111100000000000000000001"),
                                       (50002, "10.450,00", date(2026, 10, 13), "2222200000000000000000002")]:
        _escrever(pasta / f"NF {numero}.pdf",
                  pagina_danfe(numero=numero, valor=valor, vencimento=venc.strftime("%d/%m/%Y"), **emitente))
        _escrever(pasta / f"Boleto {numero}.pdf", pagina_boleto(*beneficiario, f"{numero}/1", venc, valor, livre))


if __name__ == "__main__":
    destino = Path(sys.argv[1] if len(sys.argv) > 1 else "dados/exemplos")
    for sub in ("email_alfa", "email_tech", "email_pdf_unico"):
        (destino / sub).mkdir(parents=True, exist_ok=True)
    gerar_danfe(destino / "email_alfa" / "NF 4587.pdf")
    gerar_boleto_alfa(destino / "email_alfa" / "Boleto NF 4587.pdf")
    gerar_nfse(destino / "email_tech" / "NFSe 2026000123.pdf")
    gerar_boletos_tech(destino / "email_tech" / "Boletos.pdf")
    gerar_nota_e_boleto_mesmo_pdf(destino / "email_pdf_unico" / "NF e boleto.pdf")
    gerar_notas_alfa_telecom(destino / "email_alfa_telecom")
    print(f"PDFs de exemplo gerados em {destino}/ (uma subpasta por e-mail)")

    capas = destino.parent / "capas_originais"
    for unidade in ("Centro", "Norte"):
        gerar_capa_exemplo(capas / f"2026 - Custos - Unidade {unidade}.xlsx")
    print(f"Capas de custos fictícias geradas em {capas}/")
