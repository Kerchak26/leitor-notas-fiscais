from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from nf_reader.capas import ConfigCapas, Regra, Unidade, atualizar_capas
from nf_reader.parser import NotaFiscal
from nf_reader.processamento import Lancamento
from tests.gerar_exemplos import gerar_capa_exemplo

CENTRO = "2026 - Custos - Unidade Centro.xlsx"
NORTE = "2026 - Custos - Unidade Norte.xlsx"


@pytest.fixture
def cfg(tmp_path):
    originais = tmp_path / "rede"
    gerar_capa_exemplo(originais / CENTRO)
    gerar_capa_exemplo(originais / NORTE)
    return ConfigCapas(
        pasta_capas=originais,
        pasta_conferencia=tmp_path / "conferencia",
        mes_referencia="vencimento",
        unidades=[
            Unidade("Centro", CENTRO, cnpjs=["11.111.111/0001-11"], enderecos=["Rua das Acácias, 100"]),
            Unidade("Norte", NORTE, cnpjs=["11.111.111/0002-00"], enderecos=["Av. Atlântica, 500"],
                    palavras_email=["norte"]),
        ],
        regras=[
            Regra(fornecedor="OMEGA TELECOMUNICACOES", aba="Gama - Telefonia Móvel"),
            Regra(fornecedor="ALFA TELECOM", unidade="Centro"),
        ],
    )


def nota(numero, emitente, valor, vencimento, endereco="RUA DAS ACACIAS, 100 - JARDIM",
         cnpj_dest=None, assunto="ENC: nota fiscal"):
    n = NotaFiscal(numero=str(numero), emitente=emitente, cnpj_emitente=f"00.000.{numero % 1000:03d}/0001-00",
                   valor_total=Decimal(valor), data_vencimento=vencimento, data_emissao=date(2026, 10, 1),
                   cnpj_destinatario=cnpj_dest, bloco_destinatario=f"DESTINATARIO\n{endereco}")
    return Lancamento(nota=n, origem_email=f"gestor@empresa.com.br | {assunto}")


def _situacoes(resultados):
    return {r.lanc.numero: r for r in resultados}


def _celula(cfg, arquivo, aba, celula):
    return load_workbook(cfg.pasta_conferencia / arquivo)[aba][celula].value


def test_duas_notas_do_mesmo_contrato_no_mes_sao_somadas(cfg):
    """Contrato com duas notas no mesmo mês: o valor do mês é a soma delas."""
    r = _situacoes(atualizar_capas(cfg, [
        nota(50001, "ALFA TELECOM COMERCIO E SERVICOS LTDA", "8000.00", date(2026, 10, 14)),
        nota(50002, "ALFA TELECOM COMERCIO E SERVICOS LTDA", "10450.00", date(2026, 10, 13)),
    ]))
    assert r["50001"].celula == r["50002"].celula == "M3"          # Outubro/2026
    assert _celula(cfg, CENTRO, "Alfa Telecom", "M3") == "=8000.0+10450.0"
    anotacoes = [c.value for c in load_workbook(cfg.pasta_conferencia / CENTRO)["Alfa Telecom"]["B"] if c.value]
    assert "Nota 50001 - Vencimento 14/10/2026 - R$ 8.000,00" in anotacoes
    assert "Nota 50002 - Vencimento 13/10/2026 - R$ 10.450,00" in anotacoes


def test_original_nunca_e_alterada(cfg):
    antes = (cfg.pasta_capas / CENTRO).read_bytes()
    atualizar_capas(cfg, [nota(50001, "ALFA TELECOM LTDA", "18450.00", date(2026, 10, 14))])
    assert (cfg.pasta_capas / CENTRO).read_bytes() == antes
    assert (cfg.pasta_conferencia / CENTRO).exists()


def test_nota_ja_anotada_na_capa_nao_e_lancada_de_novo(cfg):
    r = _situacoes(atualizar_capas(cfg, [nota(10101, "ALFA TELECOM LTDA", "9225.00", date(2026, 9, 14))]))
    assert r["10101"].situacao == "já está na capa original"


def test_regra_de_aba_e_unidade_por_endereco(cfg):
    r = _situacoes(atualizar_capas(cfg, [
        nota(7001, "OMEGA TELECOMUNICACOES S.A.", "2300.00", date(2026, 10, 20)),        # regra -> aba Gama
        nota(7002, "BETA LINK TELECOMUNICACOES SA", "6200.00", date(2026, 10, 10)),        # nome da aba
    ]))
    assert (r["7001"].aba, r["7001"].celula, r["7001"].criterio) == ("Gama - Telefonia Móvel", "M3", "endereço")
    assert (r["7002"].aba, r["7002"].celula) == ("Beta Link", "M3")


def test_unidade_por_cnpj_tem_prioridade_sobre_endereco(cfg):
    r = _situacoes(atualizar_capas(cfg, [
        nota(7003, "BETA LINK SA", "6200.00", date(2026, 10, 10), cnpj_dest="11.111.111/0002-00"),
    ]))
    assert (r["7003"].unidade, r["7003"].criterio) == ("Norte", "CNPJ do destinatário")


def test_contrato_escolhido_pelo_historico_de_valores(cfg):
    r = _situacoes(atualizar_capas(cfg, [
        nota(8001, "DELTA PRINT LOCACAO E SERVICOS SA", "640.00", date(2026, 10, 5)),
        nota(8002, "DELTA PRINT LOCACAO E SERVICOS SA", "1234.00", date(2026, 10, 5)),
        nota(8003, "SIGMA SISTEMAS S.A.", "750.00", date(2026, 10, 5)),
    ]))
    assert (r["8001"].aba, r["8001"].celula) == ("Delta Print", "M4")              # franquia de cópias
    assert r["8002"].situacao.startswith("pendente: aba tem 2 contratos")
    assert (r["8003"].aba, r["8003"].celula) == ("Sigma - Assinatura Digital", "M3")


def test_pendencias(cfg):
    r = _situacoes(atualizar_capas(cfg, [
        nota(9001, "BETA LINK SA", "6200.00", date(2026, 10, 10), endereco="RUA DESCONHECIDA, 1"),
        nota(9002, "FORNECEDOR NOVO LTDA", "100.00", date(2026, 10, 10)),
        nota(9003, "BETA LINK SA", "6200.00", date(2027, 1, 10)),
    ]))
    assert "unidade não identificada" in r["9001"].situacao
    assert "nenhuma aba com o nome do fornecedor" in r["9002"].situacao
    assert "não tem coluna para Janeiro/2027" in r["9003"].situacao
    relatorio = load_workbook(cfg.pasta_conferencia / "relatorio_capas.xlsx").active
    assert relatorio.max_row == 4


def test_rodar_de_novo_nao_duplica_e_original_atualizada_limpa_a_copia(cfg):
    alfa = [nota(50001, "ALFA TELECOM LTDA", "8000.00", date(2026, 10, 14)),
            nota(50002, "ALFA TELECOM LTDA", "10450.00", date(2026, 10, 13))]
    atualizar_capas(cfg, alfa)
    atualizar_capas(cfg, alfa)                       # mesmas notas de novo
    atualizar_capas(cfg, [])                         # só refaz a partir do registro
    assert _celula(cfg, CENTRO, "Alfa Telecom", "M3") == "=8000.0+10450.0"

    # Depois de conferir, o usuário copia a cópia por cima da original
    (cfg.pasta_capas / CENTRO).write_bytes((cfg.pasta_conferencia / CENTRO).read_bytes())
    resultados = atualizar_capas(cfg, [])
    assert all(r.situacao == "já está na capa original" for r in resultados)
    assert not (cfg.pasta_conferencia / CENTRO).exists()


def test_regra_nova_vale_para_pendentes_antigas(cfg):
    pendente = nota(9002, "FORNECEDOR NOVO LTDA", "750.00", date(2026, 10, 10))
    [r] = atualizar_capas(cfg, [pendente])
    assert r.situacao.startswith("pendente")
    cfg.regras.append(Regra(fornecedor="FORNECEDOR NOVO", aba="Sigma - Assinatura Digital"))
    [r] = atualizar_capas(cfg, [])                   # sem ler e-mail de novo
    assert (r.aba, r.celula) == ("Sigma - Assinatura Digital", "M3")
