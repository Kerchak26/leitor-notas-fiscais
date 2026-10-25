"""Gravação dos lançamentos (nota + boleto) em planilha Excel ou CSV, sem duplicar."""

from __future__ import annotations

import csv
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .processamento import Lancamento

COLUNAS = [
    ("numero", "Número NF"),
    ("serie", "Série"),
    ("emitente", "Emitente"),
    ("cnpj_emitente", "CNPJ Emitente"),
    ("data_emissao", "Data Emissão"),
    ("valor_nota", "Valor NF (R$)"),
    ("vencimento", "Vencimento"),
    ("valor_boleto", "Valor Boleto (R$)"),
    ("parcela", "Parcela"),
    ("linha_digitavel", "Linha Digitável"),
    ("beneficiario", "Beneficiário"),
    ("chave_acesso", "Chave de Acesso"),
    ("arquivos", "Arquivos"),
    ("origem_email", "E-mail de Origem"),
    ("avisos", "Avisos"),
    ("id", "ID"),  # usado para não gravar o mesmo lançamento duas vezes
]
TITULO_ID = COLUNAS[-1][1]


def salvar(lancamentos: list[Lancamento], destino: Path) -> int:
    """Acrescenta os lançamentos ao arquivo (xlsx ou csv). Retorna quantos eram novos."""
    destino.parent.mkdir(parents=True, exist_ok=True)
    if destino.suffix.lower() == ".csv":
        return _salvar_csv(lancamentos, destino)
    return _salvar_xlsx(lancamentos, destino)


def _novos(lancamentos: list[Lancamento], existentes: set[str]) -> list[dict]:
    resultado = []
    for lanc in lancamentos:
        d = lanc.como_dict()
        if d["id"] not in existentes:
            existentes.add(d["id"])
            resultado.append(d)
    return resultado


def _salvar_xlsx(lancamentos: list[Lancamento], destino: Path) -> int:
    if destino.exists():
        wb = load_workbook(destino)
        ws = wb.active
        col_id = [c.value for c in ws[1]].index(TITULO_ID)
        existentes = {str(r[col_id]) for r in ws.iter_rows(min_row=2, values_only=True)}
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Notas e Boletos"
        ws.append([t for _, t in COLUNAS])
        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        ws.column_dimensions[get_column_letter(len(COLUNAS))].hidden = True
        existentes = set()

    novos = _novos(lancamentos, existentes)
    for d in novos:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in (d[c] for c, _ in COLUNAS)])
        for i, (campo, _) in enumerate(COLUNAS, start=1):
            celula = ws.cell(ws.max_row, i)
            if isinstance(d[campo], date):
                celula.number_format = "DD/MM/YYYY"
            elif isinstance(d[campo], Decimal):
                celula.number_format = "#,##0.00"

    for i, (_, titulo) in enumerate(COLUNAS[:-1], start=1):
        letra = get_column_letter(i)
        largura = max(len(titulo), *(len(str(c.value or "")) for c in ws[letra]))
        ws.column_dimensions[letra].width = min(largura + 2, 60)
    wb.save(destino)
    return len(novos)


def _salvar_csv(lancamentos: list[Lancamento], destino: Path) -> int:
    existentes: set[str] = set()
    if destino.exists():
        with destino.open(encoding="utf-8-sig", newline="") as f:
            existentes = {l[TITULO_ID] for l in csv.DictReader(f, delimiter=";")}
    novos = _novos(lancamentos, existentes)
    novo_arquivo = not destino.exists()
    with destino.open("a", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")  # ";" abre direto no Excel em pt-BR
        if novo_arquivo:
            w.writerow([t for _, t in COLUNAS])
        for d in novos:
            w.writerow([_formatar_csv(d[c]) for c, _ in COLUNAS])
    return len(novos)


def _formatar_csv(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, Decimal):
        return f"{valor:.2f}".replace(".", ",")
    return str(valor)


# ---------------------------------------- controle de e-mails processados ----

def carregar_processados(arquivo: Path) -> set[str]:
    return set(json.loads(arquivo.read_text())) if arquivo.exists() else set()


def salvar_processados(ids: set[str], arquivo: Path) -> None:
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    arquivo.write_text(json.dumps(sorted(ids), indent=2))
