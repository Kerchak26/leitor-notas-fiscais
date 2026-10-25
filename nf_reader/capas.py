"""Lançamento das notas nas capas de custos (um Excel por unidade, uma aba por empresa).

A capa original nunca é alterada: o programa grava uma cópia para conferência com
tudo o que ainda não está na original. Todos os lançamentos lidos ficam registrados
em `lancamentos_capas.json`; a cada execução a cópia é refeita a partir da original
mais os lançamentos que ainda não constam nela (a anotação "Nota X" é o que indica
que uma nota já foi lançada). Assim é possível rodar várias vezes antes de conferir,
e uma regra nova em `capas.toml` vale também para as notas que ficaram pendentes.
"""

from __future__ import annotations

import json
import logging
import re
import tomllib
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .parser import data_br, normalizar
from .processamento import Lancamento

log = logging.getLogger(__name__)

MESES = ["JANEIRO", "FEVEREIRO", "MARCO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
PALAVRAS_GENERICAS = {"LTDA", "SA", "S", "A", "ME", "EPP", "EIRELI", "DE", "DA", "DO", "DOS", "DAS", "E",
                      "COMERCIO", "SERVICOS", "INDUSTRIA", "BRASIL", "DO BRASIL", "CIA", "COMPANHIA"}
COR_LANCADO = PatternFill("solid", fgColor="C6EFCE")   # verde claro: lançado agora
RE_SOMA = re.compile(r"^=\s*\d+(?:\.\d+)?(?:\s*\+\s*\d+(?:\.\d+)?)*\s*$")


# -------------------------------------------------------------- configuração ----

@dataclass
class Unidade:
    nome: str
    arquivo: str
    cnpjs: list[str] = field(default_factory=list)
    enderecos: list[str] = field(default_factory=list)
    palavras_email: list[str] = field(default_factory=list)


@dataclass
class Regra:
    fornecedor: str                       # parte do nome ou CNPJ do emitente
    unidade: str | None = None
    aba: str | None = None
    contrato: str | None = None           # texto da coluna "Contratos"
    endereco: str | None = None           # só vale se o endereço da nota contiver isto
    palavra_email: str | None = None      # só vale se o assunto do e-mail contiver isto


@dataclass
class ConfigCapas:
    pasta_capas: Path
    pasta_conferencia: Path
    mes_referencia: str
    unidades: list[Unidade]
    regras: list[Regra]

    @classmethod
    def carregar(cls, arquivo: Path) -> "ConfigCapas":
        dados = tomllib.loads(arquivo.read_text(encoding="utf-8"))
        mes = dados.get("mes_referencia", "vencimento")
        if mes not in ("vencimento", "emissao"):
            raise ValueError("mes_referencia deve ser 'vencimento' ou 'emissao'")
        return cls(
            pasta_capas=Path(dados["pasta_capas"]),
            pasta_conferencia=Path(dados.get("pasta_conferencia", "dados/capas_para_conferencia")),
            mes_referencia=mes,
            unidades=[Unidade(**u) for u in dados.get("unidades", [])],
            regras=[Regra(**r) for r in dados.get("regras", [])],
        )


# ------------------------------------------------------- lançamento salvo ----

@dataclass
class LancamentoCapa:
    """O que é preciso de uma nota para lançá-la na capa (salvo em JSON)."""

    id: str
    numero: str | None
    parcela: str | None
    emitente: str | None
    cnpj_emitente: str | None
    cnpj_destinatario: str | None
    bloco_destinatario: str | None
    origem_email: str | None
    data_emissao: str | None
    vencimento: str | None
    valor: str | None
    arquivos: str | None

    @classmethod
    def de(cls, lanc: Lancamento) -> "LancamentoCapa":
        n, b = lanc.nota, lanc.boleto
        if lanc.parcela and b and b.valor is not None:
            valor = b.valor        # nota parcelada: cada boleto no mês do seu vencimento
        else:
            valor = n.valor_total if n and n.valor_total is not None else (b.valor if b else None)
        return cls(
            id=lanc.id,
            numero=n.numero if n else None,
            parcela=lanc.parcela,
            emitente=(n.emitente if n else None) or (b.beneficiario if b else None),
            cnpj_emitente=n.cnpj_emitente if n else None,
            cnpj_destinatario=n.cnpj_destinatario if n else None,
            bloco_destinatario=n.bloco_destinatario if n else None,
            origem_email=lanc.origem_email,
            data_emissao=n.data_emissao.isoformat() if n and n.data_emissao else None,
            vencimento=lanc.vencimento.isoformat() if lanc.vencimento else None,
            valor=str(valor) if valor is not None else None,
            arquivos=lanc.como_dict()["arquivos"],
        )

    @property
    def rotulo(self) -> str:
        """Texto da anotação, no mesmo formato usado nas capas."""
        parcela = f" ({self.parcela})" if self.parcela else ""
        venc = data_br(date.fromisoformat(self.vencimento)) if self.vencimento else "sem vencimento"
        return f"Nota {self.numero}{parcela} - Vencimento {venc} - R$ {_moeda(Decimal(self.valor))}"


def carregar_registro(arquivo: Path) -> dict[str, LancamentoCapa]:
    if not arquivo.exists():
        return {}
    return {d["id"]: LancamentoCapa(**d) for d in json.loads(arquivo.read_text(encoding="utf-8"))}


def salvar_registro(registro: dict[str, LancamentoCapa], arquivo: Path) -> None:
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    dados = [asdict(l) for l in registro.values()]
    arquivo.write_text(json.dumps(dados, indent=2, ensure_ascii=False), encoding="utf-8")


# ---------------------------------------------------------------- utilidades ----

def _moeda(valor: Decimal) -> str:
    return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _compacto(texto: str) -> str:
    """Normaliza endereços/nomes para comparação: sem acento, pontuação ou abreviações."""
    t = normalizar(texto).replace("\n", " ")
    t = re.sub(r"\bAVENIDA\b", "AV", t)
    t = re.sub(r"\bRUA\b", "R", t)
    t = re.sub(r"\bNUMERO\b|\bN°|\bNO\b", "", t)
    return re.sub(r"[^A-Z0-9]", "", t)


def _palavras(texto: str) -> set[str]:
    return {p for p in re.split(r"[^A-Z0-9]+", normalizar(texto)) if p and p not in PALAVRAS_GENERICAS}


def _so_digitos(texto: str | None) -> str:
    return re.sub(r"\D", "", texto or "")


def _fornecedor_confere(regra: Regra, lanc: LancamentoCapa) -> bool:
    alvo = regra.fornecedor
    if _so_digitos(alvo) and len(_so_digitos(alvo)) >= 8:  # CNPJ (completo ou raiz)
        return _so_digitos(lanc.cnpj_emitente).startswith(_so_digitos(alvo))
    return _compacto(alvo) in _compacto(lanc.emitente or "")


def _regra_aplicavel(regras: list[Regra], lanc: LancamentoCapa) -> Regra | None:
    for regra in regras:
        if not _fornecedor_confere(regra, lanc):
            continue
        if regra.endereco and _compacto(regra.endereco) not in _compacto(lanc.bloco_destinatario or ""):
            continue
        if regra.palavra_email and normalizar(regra.palavra_email) not in normalizar(lanc.origem_email or ""):
            continue
        return regra
    return None


# ------------------------------------------------------------------ unidade ----

def identificar_unidade(cfg: ConfigCapas, lanc: LancamentoCapa, regra: Regra | None) -> tuple[Unidade | None, str]:
    """Retorna a unidade e o critério usado (ou o motivo de não ter achado)."""
    por_nome = {normalizar(u.nome): u for u in cfg.unidades}
    if regra and regra.unidade:
        unidade = por_nome.get(normalizar(regra.unidade))
        return (unidade, "regra") if unidade else (None, f"regra cita unidade inexistente '{regra.unidade}'")

    criterios = [
        ("CNPJ do destinatário", lambda u: _so_digitos(lanc.cnpj_destinatario) in {_so_digitos(c) for c in u.cnpjs}),
        ("endereço", lambda u: any(_compacto(e) in _compacto(lanc.bloco_destinatario or "") for e in u.enderecos)),
        ("assunto do e-mail", lambda u: any(normalizar(p) in normalizar(lanc.origem_email or "") for p in u.palavras_email)),
    ]
    ambiguas: list[str] = []
    for nome, confere in criterios:
        achadas = [u for u in cfg.unidades if confere(u)]
        if len(achadas) == 1:
            return achadas[0], nome
        if len(achadas) > 1:
            ambiguas = [u.nome for u in achadas]
    if ambiguas:
        return None, f"unidade ambígua ({', '.join(ambiguas)}); crie uma regra em capas.toml"
    return None, "unidade não identificada; cadastre o CNPJ/endereço ou crie uma regra"


# -------------------------------------------------------------- leitura da aba ----

@dataclass
class EstruturaAba:
    linha_cabecalho: int
    col_contratos: int
    colunas_mes: dict[tuple[int, int], int]      # (ano, mês) -> coluna
    linhas_contrato: list[int]
    linha_total: int | None


def _valor_mesclado(ws: Worksheet, linha: int, coluna: int):
    for faixa in ws.merged_cells.ranges:
        if faixa.min_row <= linha <= faixa.max_row and faixa.min_col <= coluna <= faixa.max_col:
            return ws.cell(faixa.min_row, faixa.min_col).value
    return ws.cell(linha, coluna).value


def _mes_do_cabecalho(valor) -> tuple[int | None, int] | None:
    if isinstance(valor, (datetime, date)):
        return valor.year, valor.month
    if not isinstance(valor, str):
        return None
    texto = normalizar(valor).strip()
    for i, mes in enumerate(MESES, start=1):
        if texto.startswith(mes) or re.match(rf"^{mes[:3]}\b", texto):
            ano = re.search(r"(20\d{2})", texto) or re.search(r"/(\d{2})\b", texto)
            if ano:
                a = int(ano.group(1))
                return (a + 2000 if a < 100 else a), i
            return None, i
    return None


def ler_estrutura(ws: Worksheet, ano_padrao: int) -> EstruturaAba | None:
    max_col = min(ws.max_column, 60)
    for linha in range(1, min(ws.max_row, 15) + 1):
        meses = {c: _mes_do_cabecalho(ws.cell(linha, c).value) for c in range(1, max_col + 1)}
        meses = {c: m for c, m in meses.items() if m}
        if len(meses) >= 3:
            break
    else:
        return None

    # Ano: no próprio cabeçalho, nas linhas de cima (célula mesclada "2026") ou por sequência
    colunas_mes: dict[tuple[int, int], int] = {}
    ano_corrente, mes_anterior = None, None
    for coluna in sorted(meses):
        ano, mes = meses[coluna]
        if ano is None:
            for acima in range(linha - 1, 0, -1):
                v = _valor_mesclado(ws, acima, coluna)
                if isinstance(v, (int, float)) and 2000 <= v <= 2100 or isinstance(v, str) and re.fullmatch(r"\s*20\d{2}\s*", v):
                    ano = int(v)
                    break
        if ano is None:
            if ano_corrente is None:
                ano = ano_padrao
            else:
                ano = ano_corrente + (1 if mes_anterior and mes < mes_anterior else 0)
        colunas_mes[(ano, mes)] = coluna
        ano_corrente, mes_anterior = ano, mes

    primeira_mes = min(meses)
    col_contratos = next((c for c in range(1, primeira_mes)
                          if "CONTRATO" in normalizar(str(ws.cell(linha, c).value or ""))), primeira_mes - 1)
    linhas_contrato, linha_total = [], None
    for l in range(linha + 1, ws.max_row + 1):
        textos = " ".join(normalizar(str(ws.cell(l, c).value or "")) for c in range(1, primeira_mes))
        if "TOTAL" in textos:
            linha_total = l
            break
        if ws.cell(l, col_contratos).value not in (None, ""):
            linhas_contrato.append(l)
    return EstruturaAba(linha, col_contratos, colunas_mes, linhas_contrato, linha_total)


# ----------------------------------------------------------- aba e contrato ----

def _valores_da_linha(ws: Worksheet, linha: int, estrutura: EstruturaAba) -> set[Decimal]:
    valores = set()
    for coluna in estrutura.colunas_mes.values():
        v = ws.cell(linha, coluna).value
        if isinstance(v, (int, float)) and v:
            valores.add(Decimal(str(v)).quantize(Decimal("0.01")))
    return valores


def escolher_aba(wb: Workbook, lanc: LancamentoCapa, regra: Regra | None) -> tuple[list[str], str]:
    """Abas candidatas para o fornecedor da nota."""
    if regra and regra.aba:
        nomes = [n for n in wb.sheetnames if normalizar(n) == normalizar(regra.aba)]
        return nomes, "" if nomes else f"aba '{regra.aba}' (da regra) não existe na capa"
    emitente = _palavras(lanc.emitente or "")
    emitente_junto = _compacto(lanc.emitente or "")
    candidatas = []
    for nome in wb.sheetnames:
        chave = re.split(r"\s+-|-\s+|\(", nome)[0]       # "Sigma - Assinatura Digital" -> "Sigma"
        palavras = _palavras(chave)
        junto = _compacto(chave)
        # Nomes curtos ("ZE", "LUZ") só por palavra inteira, para não achar "AZEVEDO" ou "LUZIA"
        if palavras and (palavras <= emitente or (len(junto) >= 5 and junto in emitente_junto)):
            candidatas.append(nome)
    if not candidatas:
        return [], f"nenhuma aba com o nome do fornecedor '{lanc.emitente}'; crie uma regra com a aba"
    return candidatas, ""


def escolher_linha(ws: Worksheet, estrutura: EstruturaAba, valor: Decimal, regra: Regra | None) -> tuple[int | None, str]:
    linhas = estrutura.linhas_contrato
    if regra and regra.contrato:
        achadas = [l for l in linhas if _compacto(regra.contrato) in _compacto(str(ws.cell(l, estrutura.col_contratos).value))]
        return (achadas[0], "") if achadas else (None, f"contrato '{regra.contrato}' não encontrado na aba")
    if len(linhas) == 1:
        return linhas[0], ""
    if not linhas:
        return None, "aba sem linhas de contrato"
    # Contratos têm valor recorrente: procura a linha que já teve este valor em outro mês
    iguais = [l for l in linhas if valor in _valores_da_linha(ws, l, estrutura)]
    if len(iguais) == 1:
        return iguais[0], ""
    nomes = [str(ws.cell(l, estrutura.col_contratos).value) for l in linhas]
    return None, f"aba tem {len(linhas)} contratos ({'; '.join(nomes)}); indique o contrato numa regra"


# ------------------------------------------------------------------ escrita ----

def _ja_lancada(ws: Worksheet, numero: str, parcela: str | None) -> bool:
    padrao = re.compile(rf"\bNOTA\s*(?:N[°O.]*\s*)?:?\s*0*{re.escape(numero)}\b"
                        + (rf"\s*\({re.escape(parcela)}\)" if parcela else r"(?!\s*\(\d)"))
    for linha in ws.iter_rows(values_only=True):
        for v in linha:
            if isinstance(v, str) and padrao.search(normalizar(v)):
                return True
    return False


def _ultima_linha_usada(ws: Worksheet) -> int:
    for linha in range(ws.max_row, 0, -1):
        if any(c.value not in (None, "") for c in ws[linha]):
            return linha
    return 0


def _escrever_anotacao(ws: Worksheet, coluna: int, texto: str) -> None:
    linha = _ultima_linha_usada(ws) + 1
    mescladas = {(l, c) for faixa in ws.merged_cells.ranges
                 for l in range(faixa.min_row, faixa.max_row + 1) for c in range(faixa.min_col, faixa.max_col + 1)}
    while (linha, coluna) in mescladas:
        linha += 1
    ws.cell(linha, coluna, texto).fill = COR_LANCADO


def lancar_valor(ws: Worksheet, linha: int, coluna: int, valor: Decimal, rotulo: str) -> tuple[bool, str]:
    celula = ws.cell(linha, coluna)
    atual = celula.value
    numero = float(valor)
    if atual in (None, ""):
        celula.value = numero
        situacao = "lançado"
    elif isinstance(atual, (int, float)):
        celula.value = f"={atual}+{numero}"
        situacao = f"somado ao valor que já estava na célula (R$ {_moeda(Decimal(str(atual)))})"
    elif isinstance(atual, str) and RE_SOMA.match(atual):
        celula.value = f"{atual}+{numero}"
        situacao = "somado às outras notas do mês"
    else:
        return False, f"célula {celula.coordinate} tem conteúdo que não é número ({atual!r}); lance manualmente"

    # Mesmo formato numérico do mês ao lado
    for vizinha in (coluna - 1, coluna + 1):
        fmt = ws.cell(linha, vizinha).number_format
        if fmt and fmt != "General":
            celula.number_format = fmt
            break
    else:
        celula.number_format = "#,##0.00"
    celula.fill = COR_LANCADO
    texto = (celula.comment.text + "\n" if celula.comment else "") + rotulo
    celula.comment = Comment(texto, "nf_reader", width=320, height=60 + 20 * texto.count("\n"))
    return True, f"{situacao} em {celula.coordinate}"


# ---------------------------------------------------------------- principal ----

@dataclass
class Resultado:
    lanc: LancamentoCapa
    situacao: str
    unidade: str = ""
    aba: str = ""
    celula: str = ""
    criterio: str = ""


def _ano_do_arquivo(nome: str) -> int:
    m = re.search(r"(20\d{2})", nome)
    return int(m.group(1)) if m else date.today().year


def atualizar_capas(cfg: ConfigCapas, novos: list[Lancamento]) -> list[Resultado]:
    arquivo_registro = cfg.pasta_conferencia / "lancamentos_capas.json"
    registro = carregar_registro(arquivo_registro)
    for lanc in novos:
        if lanc.nota or lanc.boleto:
            registro.setdefault(lanc.id, LancamentoCapa.de(lanc))
    salvar_registro(registro, arquivo_registro)

    resultados: list[Resultado] = []
    por_unidade: dict[str, list[tuple[LancamentoCapa, Regra | None, str]]] = {}
    for lanc in registro.values():
        if not lanc.numero or lanc.valor is None:
            resultados.append(Resultado(lanc, "pendente: nota sem número ou sem valor"))
            continue
        data_ref = lanc.vencimento if cfg.mes_referencia == "vencimento" else lanc.data_emissao
        if not data_ref:
            resultados.append(Resultado(lanc, f"pendente: nota sem data de {cfg.mes_referencia}"))
            continue
        regra = _regra_aplicavel(cfg.regras, lanc)
        unidade, criterio = identificar_unidade(cfg, lanc, regra)
        if not unidade:
            resultados.append(Resultado(lanc, f"pendente: {criterio}"))
            continue
        por_unidade.setdefault(unidade.arquivo, []).append((lanc, regra, criterio))

    for arquivo, itens in por_unidade.items():
        resultados += _atualizar_arquivo(cfg, arquivo, itens)
    _salvar_relatorio(resultados, cfg.pasta_conferencia / "relatorio_capas.xlsx")
    return resultados


def _atualizar_arquivo(cfg: ConfigCapas, arquivo: str, itens) -> list[Resultado]:
    original = cfg.pasta_capas / arquivo
    unidade_nome = next(u.nome for u in cfg.unidades if u.arquivo == arquivo)
    if not original.exists():
        return [Resultado(l, f"pendente: capa não encontrada ({original})", unidade_nome) for l, _, _ in itens]

    wb = load_workbook(original, keep_vba=original.suffix.lower() == ".xlsm")
    estruturas: dict[str, EstruturaAba | None] = {}
    resultados, alterou = [], False

    for lanc, regra, criterio in itens:
        r = Resultado(lanc, "", unidade_nome, criterio=criterio)
        resultados.append(r)
        valor = Decimal(lanc.valor)
        candidatas, motivo = escolher_aba(wb, lanc, regra)
        if not candidatas:
            r.situacao = f"pendente: {motivo}"
            continue

        # Já lançada (pela anotação "Nota X") em alguma das abas candidatas?
        if any(_ja_lancada(wb[a], lanc.numero, lanc.parcela) for a in candidatas):
            r.situacao, r.aba = "já está na capa original", candidatas[0]
            continue

        destinos = []
        for aba in candidatas:
            if aba not in estruturas:
                estruturas[aba] = ler_estrutura(wb[aba], _ano_do_arquivo(arquivo))
            if estruturas[aba]:
                linha, motivo_linha = escolher_linha(wb[aba], estruturas[aba], valor, regra)
                destinos.append((aba, linha, motivo_linha))
        validos = [d for d in destinos if d[1]]
        if len(validos) > 1:
            # Várias abas possíveis (ex.: "Sigma" e "Sigma - Assinatura Digital"): usa o histórico de valores
            pelo_valor = [d for d in validos if valor in _valores_da_linha(wb[d[0]], d[1], estruturas[d[0]])]
            if len(pelo_valor) == 1:
                validos = pelo_valor
        if len(validos) != 1:
            if len(candidatas) > 1 and not validos:
                r.situacao = f"pendente: mais de uma aba possível ({', '.join(candidatas)}); crie uma regra"
            elif len(validos) > 1:
                r.situacao = f"pendente: mais de uma aba possível ({', '.join(d[0] for d in validos)}); crie uma regra"
            else:
                r.situacao = "pendente: " + (destinos[0][2] if destinos else f"não reconheci o cabeçalho de meses da aba {candidatas[0]}")
            continue

        aba, linha, _ = validos[0]
        ws, estrutura = wb[aba], estruturas[aba]
        r.aba = aba
        data_ref = date.fromisoformat(lanc.vencimento if cfg.mes_referencia == "vencimento" else lanc.data_emissao)
        coluna = estrutura.colunas_mes.get((data_ref.year, data_ref.month))
        if not coluna:
            r.situacao = f"pendente: a aba não tem coluna para {MESES[data_ref.month - 1].title()}/{data_ref.year}"
            continue
        ok, situacao = lancar_valor(ws, linha, coluna, valor, lanc.rotulo)
        r.situacao = situacao if ok else f"pendente: {situacao}"
        if ok:
            r.celula = ws.cell(linha, coluna).coordinate
            _escrever_anotacao(ws, estrutura.col_contratos, lanc.rotulo)
            alterou = True

    destino = cfg.pasta_conferencia / arquivo
    if alterou:
        destino.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destino)
        log.info("Cópia para conferência gravada: %s", destino)
    elif destino.exists():
        destino.unlink()  # tudo já está na original; a cópia antiga só confundiria
        log.info("Nada novo para %s; cópia antiga removida", arquivo)
    return resultados


def _salvar_relatorio(resultados: list[Resultado], destino: Path) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    titulos = ["Situação", "Unidade", "Aba", "Célula", "Nota", "Emitente", "Vencimento", "Valor (R$)",
               "Unidade identificada por", "E-mail", "Arquivos"]
    ws.append(titulos)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    ordem = sorted(resultados, key=lambda r: (not r.situacao.startswith("pendente"), r.unidade, r.aba))
    for r in ordem:
        l = r.lanc
        ws.append([
            r.situacao, r.unidade, r.aba, r.celula,
            f"{l.numero}" + (f" ({l.parcela})" if l.parcela else "") if l.numero else "",
            l.emitente, date.fromisoformat(l.vencimento) if l.vencimento else None,
            float(l.valor) if l.valor else None, r.criterio, l.origem_email, l.arquivos,
        ])
        ws.cell(ws.max_row, 7).number_format = "DD/MM/YYYY"
        ws.cell(ws.max_row, 8).number_format = "#,##0.00"
        if r.situacao.startswith("pendente"):
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor="FFC7CE")
        elif not r.situacao.startswith("já"):
            ws.cell(ws.max_row, 1).fill = COR_LANCADO
    for i, largura in enumerate([60, 14, 22, 8, 12, 36, 12, 14, 22, 40, 50], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = largura
    ws.freeze_panes = "A2"
    wb.save(destino)
